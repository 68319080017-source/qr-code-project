import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from io import BytesIO
from typing import List, Dict, Any
from datetime import datetime

def export_assets_to_excel(assets: List[Any]) -> BytesIO:
    """
    Exports a list of assets to an Excel file.
    Returns a BytesIO stream containing the Excel file.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assets"
    
    # Define headers
    headers = [
        "Asset Number", "Asset Code", "Name", "Category", "Status", 
        "Location", "Responsible", "Purchase Date", "Price"
    ]
    
    # Style for headers
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    
    # Write headers
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 20
        
    # Write data
    for row_num, asset in enumerate(assets, 2):
        ws.cell(row=row_num, column=1, value=asset.asset_number)
        ws.cell(row=row_num, column=2, value=asset.asset_code or "-")
        ws.cell(row=row_num, column=3, value=asset.name)
        ws.cell(row=row_num, column=4, value=asset.category.name if asset.category else "-")
        ws.cell(row=row_num, column=5, value=asset.status)
        
        location_str = f"{asset.location.building} {asset.location.room}" if asset.location else "-"
        ws.cell(row=row_num, column=6, value=location_str)
        
        responsible = f"{asset.responsible_person.first_name} {asset.responsible_person.last_name}" if asset.responsible_person else "-"
        ws.cell(row=row_num, column=7, value=responsible)
        
        purchase_date = asset.purchase_date.strftime("%Y-%m-%d") if asset.purchase_date else "-"
        ws.cell(row=row_num, column=8, value=purchase_date)
        
        ws.cell(row=row_num, column=9, value=float(asset.price) if asset.price else 0.0)
        
    # Save to stream
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return output

def import_assets_from_excel(file_content: bytes) -> List[Dict[str, Any]]:
    """
    Parses an uploaded Excel file to extract asset data.
    """
    wb = openpyxl.load_workbook(filename=BytesIO(file_content))
    ws = wb.active
    
    assets_data = []
    # Skip header row, assuming it's row 1
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]: # Skip empty rows
            continue
            
        assets_data.append({
            "asset_number": str(row[0]),
            "name": str(row[2]) if len(row) > 2 and row[2] else "Imported Asset",
            # Add other fields mapping based on template
        })
        
    return assets_data
